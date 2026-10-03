"""
Builds the manuscript METHODS section (machine-learning component) as a
standalone, submission-formatted .docx, condensed to fit an 8-page limit and
referenced throughout.

Every quantity is pulled from outputs/metrics/* at build time, and every
citation marker resolves against 06_Code/references.py, whose entries were each
verified against PubMed (or the publisher's DOI record) before use.

Numbering convention, chosen so this file and the Results file can be assembled
without collisions: Figure 1 is the participant-flow diagram, which the Results
also cites as Figure 1 because it is the same figure. All other Methods exhibits
are supplementary (Table S1-S3, Figure S1), which is where journals normally
place pipeline schematics and variable dictionaries anyway.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx_style as ds
from references import REFS, cite

ROOT = Path(__file__).resolve().parents[1]
MET = ROOT / "outputs" / "metrics"
FIGP = ROOT / "outputs" / "figures_pub"
OUT = ROOT / "manuscript"
OUT.mkdir(parents=True, exist_ok=True)

N = json.loads((MET / "publication_numbers.json").read_text())
tune = pd.read_csv(MET / "tuning_summary.csv")
miss = pd.read_csv(MET / "missingness_report.csv")
coll = pd.read_csv(MET / "collinearity_flags.csv")
outdef = pd.read_csv(MET / "outcome_definition_comparison.csv")

f = N["flow"]
n_train, n_test = f["train_n"], f["test_n"]
n_pred = N["n_predictors"]
n_cols = N["n_transformed_features"]
events = round(N["train_prevalence_pct"] / 100 * n_train)
NHANES_URL = "https://wwwn.cdc.gov/nchs/nhanes/"


def main():
    doc = ds.new_document(line_spacing="1.5")

    ds.title(doc, "Methods")
    ds.subtitle(doc, "Machine-learning analysis of blood metals and metabolic syndrome "
                     "in NHANES 2007–2018")

    # ==================================================================
    ds.heading(doc, "1. Study design, data source and reporting", 1)

    ds.para(doc,
            f"This cross-sectional prediction-model development study uses six consecutive "
            f"two-year cycles of the National Health and Nutrition Examination Survey (NHANES), "
            f"2007–2008 through 2017–2018. NHANES applies a complex, multistage probability design "
            f"to sample the civilian non-institutionalized United States population, pairing a "
            f"household interview with a standardized examination and laboratory assessment "
            f"{cite('nhanes_design', 'nhanes_weights')}. All data are public and were downloaded "
            f"from the National Center for Health Statistics ({NHANES_URL}). The thirteen survey "
            f"modules used are listed in Table S1; all were merged on the respondent sequence "
            f"number (SEQN), with the demographic file defining the sampled-person universe. "
            f"Reporting follows TRIPOD+AI {cite('tripod')} and STROBE {cite('strobe')}. The "
            f"Mendelian randomization component is reported separately.")

    # ==================================================================
    ds.heading(doc, "2. Study population and data cleaning", 1)

    ds.para(doc,
            f"Figure 1 gives the data-cleaning sequence with the number excluded at each step and "
            f"the number remaining. From {f['pooled']:,} pooled participants we excluded "
            f"{f['excluded_age']:,} aged under 20 years, {f['excluded_pregnant']:,} pregnant "
            f"participants (pregnancy alters waist circumference, blood pressure, lipids and "
            f"glucose), {f['excluded_mets']:,} adults missing any measurement required to define the "
            f"outcome (mainly because triglycerides and fasting glucose are restricted to the morning "
            f"fasting subsample), {f['excluded_fasting']:,} who did not meet the fasting criterion of "
            f"8 to 24 hours (identified by a zero fasting-subsample weight, which NCHS assigns to "
            f"every such participant in every cycle), and {f['excluded_metals']:,} with no blood "
            f"metal assayed, leaving {f['final_cohort']:,} adults. Selenium and manganese were not "
            f"assayed in 2007–2008 or 2009–2010; those participants were retained because they "
            f"contribute complete lead, cadmium and mercury data, making the resulting absence "
            f"structural rather than random (Section 5). Refusal and don't-know codes were recoded "
            f"to missing using each item's own code width (for example 7 and 9 for one-digit items "
            f"but 777 and 999 for the number of drinks), so that genuine answers of 7 or 9 drinks "
            f"were kept.")

    ds.figure(doc, FIGP / "fig1_participant_flow.png", width_in=5.4)
    ds.caption(doc, "Figure 1.",
               "Data-cleaning and participant-flow diagram. Boxes on the left give the number "
               "remaining after each step; boxes on the right give the number excluded at that "
               "step and the reason. Cycle-level contributions to the pooled sample are shown "
               "above the diagram. MetS, metabolic syndrome.")

    # ==================================================================
    ds.heading(doc, "3. Outcome definition", 1)

    ds.para(doc,
            f"Metabolic syndrome was defined by the harmonized criteria of the joint interim "
            f"statement of the International Diabetes Federation, NHLBI, American Heart Association "
            f"and collaborating bodies {cite('alberti')}, which retain the AHA/NHLBI thresholds and "
            f"drug-treatment clauses {cite('grundy')}: a component was scored positive when the "
            f"measured value met the cut-point or the participant reported drug treatment for it. "
            f"Cases met at least three of five components; definitions, thresholds and source "
            f"variables are in Table S2. Blood pressure was the mean of all valid readings (up to "
            f"four; a diastolic reading of 0 mmHg was treated as invalid). NHANES asks about "
            f"cholesterol-lowering treatment only once, without naming the drug, whereas the "
            f"guideline's treatment clause for triglycerides and HDL cholesterol refers to "
            f"lipid-specific drugs such as fibrates and nicotinic acid; we applied the item to both "
            f"lipid components as a stated modelling choice and tested it in a sensitivity analysis "
            f"(Section 9). The treatment clauses together moved "
            f"{int(outdef.loc[1, 'n_reclassified_negative_to_positive']):,} participants from "
            f"non-case to case and none in reverse, raising prevalence from "
            f"{outdef.loc[0, 'prevalence_pct']:.2f}% to {outdef.loc[1, 'prevalence_pct']:.2f}%. The "
            f"component measurements and medication items build the label only and were excluded "
            f"from the predictor set, since a model given any of them could recover the outcome "
            f"arithmetically rather than learn a risk relationship {cite('bernett')}. Body mass index "
            f"is a predictor while waist circumference defines one component; the two are strongly "
            f"correlated, so part of body mass index's predictive weight is close to definitional, and "
            f"a model without it was fitted for comparison (Section 9).")

    # ==================================================================
    ds.heading(doc, "4. Predictor variables", 1)

    ds.para(doc,
            f"Predictors were specified in advance from the metabolic-syndrome and metal-toxicity "
            f"literature, across four domains. The **blood-metal panel** is lead, cadmium, total "
            f"mercury, selenium and manganese — the analytes measured together in the NHANES "
            f"whole-blood assay and the panel used in prior NHANES analyses of metals and metabolic "
            f"syndrome {cite('bulka')}. **Demographic and socioeconomic** predictors (age, sex, "
            f"race/ethnicity, education, marital status, household size, income-to-poverty ratio, "
            f"examination period) are established correlates of both the outcome and metal exposure "
            f"and are retained as adjustment variables. The **anthropometric** domain is body mass "
            f"index. **Lifestyle** predictors cover smoking, alcohol use, self-rated diet quality, "
            f"four physical-activity items and resting pulse. Table S1 gives every variable with its "
            f"NHANES item code, a plain description of what that code records, its coding and the "
            f"model columns it contributes; after encoding the set held {n_pred} columns from 24 "
            f"source variables.")

    ds.para(doc,
            f"No data-driven selection procedure — stepwise regression or the least absolute "
            f"shrinkage and selection operator — was used in this phase. Selection algorithms run on "
            f"a fixed dataset produce model-selection instability, biased effect estimates and "
            f"confidence intervals whose nominal coverage fails, so the surviving set partly reflects "
            f"a stochastic screening step rather than which variables matter {cite('heinze')}. Since "
            f"this phase exists to rank candidate exposures for the Mendelian randomization "
            f"analysis, a step that discards variables before ranking would remove the quantity of "
            f"interest. Penalized selection retains a clear role there, applied to the "
            f"high-dimensional genetic instrument space rather than a pre-specified "
            f"{n_pred}-column epidemiological set.")

    r = coll.iloc[0]
    ds.para(doc,
            f"Two candidates were removed prospectively. Hip circumference was excluded on training-"
            f"partition grounds — correlation {r['pearson_r']:.3f} with body mass index and "
            f"{r['pct_missing_var_1_train']:.1f}% missing — so it adds imputation burden while "
            f"splitting importance with a near-identical retained variable. Physician-diagnosed "
            f"diabetes was held out as partly downstream of the glucose component and reintroduced "
            f"only in sensitivity analysis. Three items needed harmonization or skip-pattern "
            f"handling. The alcohol questionnaire was redesigned in 2017–2018, and its gatekeeper "
            f"questions differ between periods (\"at least 12 drinks in any one year\" before, "
            f"\"ever had a drink\" after), so alcohol use was represented by a construct asked the "
            f"same way in every cycle: any alcohol in the past 12 months (from the past-year drinking "
            f"frequency), with average drinks per drinking day set to zero for past-year "
            f"non-drinkers, who are never asked that question. Current smoking intensity is asked "
            f"only of people who have smoked at least 100 cigarettes, so never-smokers were coded as "
            f"not smoking now. PAD645 (minutes walking or bicycling for transportation) is asked "
            f"only of those answering yes to PAQ635, so a blank after a *no* denotes zero minutes and "
            f"was recoded accordingly.")

    # ==================================================================
    ds.heading(doc, "5. Missing data", 1)

    top = miss.set_index("variable")["pct_missing"]
    preds_base = set(N["missing_flags"]) | {"ALQ_PAST12", "SMQ040", "DMDEDUC2", "DMDMARTL", "SMQ020",
                                           "PAQ605", "PAQ650", "PAQ665", "DBQ700"}
    others = top[[v for v in top.index if v in preds_base and v not in
                  ("LBXBSE", "LBXBMN", "INDFMPIR", "ALQ130", "ALQ_PAST12")]]
    ds.para(doc,
            f"After the skip patterns were resolved, missingness was concentrated in selenium and "
            f"manganese ({top['LBXBSE']:.1f}% each, entirely the two cycles lacking the assay), the "
            f"income-to-poverty ratio ({top['INDFMPIR']:.1f}%) and the two alcohol items "
            f"({top['ALQ_PAST12']:.1f}% and {top['ALQ130']:.1f}%); every other predictor was under "
            f"{max(1, int(others.max()) + 1)}%. Values were imputed inside the "
            f"pipeline using the median for continuous and the modal category for categorical "
            f"variables, fitted on training folds only. Multiple imputation is standard for "
            f"etiological analysis {cite('sterne')} but sits awkwardly with a model that must also "
            f"score a future incomplete record, and simulation work spanning development and "
            f"deployment finds single imputation with missing indicators comparable for that purpose "
            f"{cite('sisk')}. Each continuous predictor with missing values therefore carries a "
            f"binary indicator (1 = missing and imputed); selenium and manganese share one indicator "
            f"because they are missing for exactly the same participants, which makes it a marker of "
            f"the 2007–2010 cycles rather than of individual non-response. The indicators are fixed "
            f"in advance, so the transformed matrix holds a constant {n_cols} columns in every fold. "
            f"Indicators let the model use informative missingness but can bias predictions when "
            f"missingness depends on the outcome {cite('sisk')}; they are reported separately from "
            f"the substantive predictors rather than interpreted as risk factors.")

    # ==================================================================
    ds.heading(doc, "6. Partitioning and prevention of information leakage", 1)

    ds.para(doc,
            f"The cohort was split once into {n_train:,} training (80%) and {n_test:,} held-out test "
            f"participants (20%), stratified on outcome. The test partition was opened only after "
            f"development was complete and contributed to no imputation statistic, scaling "
            f"parameter, resampling step, hyperparameter choice or decision threshold. With "
            f"{events:,} events and {n_pred} candidate columns the training partition provides about "
            f"{events / n_pred:.0f} events per candidate predictor; that benchmark was derived for "
            f"regression models and does not translate directly to the effective complexity of a "
            f"boosted ensemble {cite('riley')}, so it is a descriptive adequacy check rather than a "
            f"licence for complexity, which was controlled by cross-validated tuning instead.")

    ds.para(doc,
            f"Every step that estimates anything from data sits inside one pipeline object, so "
            f"cross-validation refits the whole sequence within each training fold (Figure S1). "
            f"Fitting imputation, scaling or synthetic oversampling on the full dataset before "
            f"splitting is among the most common forms of data leakage: a systematic survey "
            f"identified 294 affected papers across 17 fields, and in several the corrected analysis "
            f"erased the reported advantage of complex models over logistic regression "
            f"{cite('kapoor')}. Leakage through preprocessing and feature handling is specifically "
            f"flagged as a recurring failure in biomedical machine learning {cite('bernett')}. The "
            f"variable screens in Section 4 were likewise computed on training data alone.")

    # ==================================================================
    ds.heading(doc, "7. Model development", 1)

    ds.para(doc,
            f"Continuous predictors were median-imputed and standardized; binary, ordinal and "
            f"one-hot columns were mode-imputed and left on their original scale. Standardization is "
            f"required because the model set includes distance- and kernel-based learners — "
            f"k-nearest neighbors, the support vector machine and the multilayer perceptron — whose "
            f"objectives are not scale-invariant, so an unscaled predictor measured in hundreds "
            f"(selenium, µg/L) would dominate the distance metric over one measured in units. Tree "
            f"ensembles are invariant to monotone rescaling and are unaffected, so one shared "
            f"preprocessing path keeps all eleven models comparable. Indicator columns are left "
            f"unscaled because standardizing a 0/1 column changes nothing for trees and makes linear "
            f"coefficients harder to read. Column order after transformation is deterministic and "
            f"identical in every fold, which is what lets SHAP values be aligned unambiguously to "
            f"predictor names.")

    ds.para(doc,
            f"Training-partition prevalence was {N['train_prevalence_pct']:.1f}%, so the classes are "
            f"near-balanced and the usual case for resampling is weak. Synthetic Minority "
            f"Over-sampling for Nominal and Continuous features (SMOTENC) was nonetheless applied "
            f"within each training fold. SMOTENC rather than SMOTE was necessary because the "
            f"predictor set mixes continuous measurements with binary and one-hot columns: SMOTE "
            f"interpolates linearly between neighbors and would place fractional values on "
            f"categorical columns, creating synthetic records with impossible category memberships, "
            f"whereas SMOTENC interpolates only continuous features and assigns each synthetic record "
            f"the majority category among its neighbors {cite('chawla')}. The countervailing "
            f"evidence is also relevant: imbalance correction can leave risk models badly "
            f"miscalibrated without improving discrimination, and shifting the decision threshold "
            f"achieves the same sensitivity–specificity trade-off without that cost "
            f"{cite('goorbergh')}. We therefore treated the choice as a parameter to test rather "
            f"than assume, compared SMOTENC against class weighting in a prespecified sensitivity "
            f"analysis (Section 9), and report calibration alongside discrimination for every model.")

    ds.para(doc,
            f"Eleven classifiers spanning linear, instance-based, kernel, neural and tree-ensemble "
            f"families were trained: logistic regression, decision tree, random forest, gradient "
            f"boosting, LightGBM, XGBoost, CatBoost, support vector machine, Gaussian naïve Bayes, "
            f"multilayer perceptron and k-nearest neighbors. Hyperparameters were tuned over the "
            f"grids in Table S3 with 10-fold stratified cross-validation on the training partition, "
            f"optimizing AUROC; every grid had at most 12 configurations and was searched "
            f"exhaustively, and the selected configurations are reported there. "
            f"Thresholds were selected by maximizing the Youden index {cite('youden')} on "
            f"out-of-fold training predictions, since choosing a threshold on the data used to report "
            f"sensitivity and specificity inflates both; results are given at the tuned threshold "
            f"and at 0.5.")

    # ==================================================================
    ds.heading(doc, "8. Performance evaluation and model interpretation", 1)

    ds.para(doc,
            f"Discrimination was summarized by AUROC and the area under the precision–recall curve, "
            f"with accuracy, sensitivity, specificity, predictive values, F1 and Cohen's kappa at "
            f"both thresholds; every test-set metric carries a 1,000-resample percentile bootstrap "
            f"interval. Because all models are scored on the same participants their AUROC estimates "
            f"are correlated, so comparing independent intervals would overstate the uncertainty of a "
            f"difference; each model was compared with the leading model by DeLong's test for "
            f"correlated AUROCs {cite('delong')}, with the interval of the difference from a paired "
            f"bootstrap applying the same resampled indices to both models in every draw. "
            f"Calibration was assessed with quantile-binned curves, the Brier score, the calibration "
            f"slope, the calibration-in-the-large intercept and the observed-to-expected ratio, "
            f"because a model can rank well while misstating absolute "
            f"risk — the failure that matters if a predicted probability reaches a clinician "
            f"{cite('vancalster')}. Clinical utility was assessed by decision-curve analysis against "
            f"treat-all and treat-none strategies over threshold probabilities from 2% to 80% "
            f"{cite('vickers')}.")

    ds.para(doc,
            f"Predictor contributions were quantified two ways. SHapley Additive exPlanations were "
            f"computed with the exact tree explainer {cite('lundberg')}, giving each predictor an "
            f"additive contribution to each prediction on the log-odds scale, with global importance "
            f"the mean absolute value. Permutation importance was the mean decrease in AUROC over ten "
            f"permutations of each column, with the standard deviation retained so rank differences "
            f"can be read against the noise producing them, in the permutation framework of "
            f"{cite('altmann')}. The two measure "
            f"different things — how far a predictor moves individual predictions, versus how much "
            f"ranking performance depends on it — so a predictor substitutable by correlated "
            f"predictors scores high on the first and low on the second, and both are reported. SHAP "
            f"runs on the {n_cols}-column transformed matrix and permutation importance on the "
            f"{n_pred} original columns, so SHAP importances are ranked among those {n_pred} "
            f"predictors to keep one denominator. Cross-model agreement was assessed by recomputing "
            f"permutation importance independently in each of the four leading models.")

    # ==================================================================
    ds.heading(doc, "9. Sensitivity and subgroup analyses", 1)

    ds.para(doc,
            f"Sensitivity analyses addressed the design choices most likely to affect conclusions: "
            f"adding physician-diagnosed diabetes; replacing SMOTENC with class weighting (balanced "
            f"sample weights for gradient boosting, which has no class-weight option); restricting "
            f"the panel to blood lead alone; and three alternative outcome definitions "
            f"(measured values only; no cholesterol-medication clause for the lipid components; "
            f"blood pressure from the first reading only). Each used 10-fold cross-validation on the "
            f"training partition with the three leading algorithms at their tuned hyperparameters, "
            f"holding everything else fixed. The incremental value of the metals was assessed by "
            f"refitting the leading model without the five metals, with lead as the only metal and "
            f"without body mass index, and by a logistic regression on age and body mass index "
            f"alone; each was scored on the same test participants and compared with the full model "
            f"by DeLong's test. Because a single random partition leaves split-to-split variability "
            f"uncaptured, all eleven models were also evaluated by repeated stratified 10-fold "
            f"cross-validation (five repeats) on the training partition. Discrimination and "
            f"calibration were "
            f"also recomputed within strata of sex, age tertile and race/ethnicity, with bootstrap "
            f"intervals and the calibration model refitted in every resample. Subgroup analyses are "
            f"descriptive; each stratum is nested within the overall sample, so subgroup intervals "
            f"are not independent of the overall estimate and are not used to test against it.")

    # ==================================================================
    ds.heading(doc, "10. Scope, software and availability", 1)

    ds.para(doc,
            f"This is a development study with internal validation: performance is estimated on a "
            f"held-out partition of the same cohort, supplemented by repeated cross-validation and a "
            f"temporal validation in which the leading algorithm was re-tuned and refitted on the "
            f"2007–2014 cycles and tested on the 2015–2018 cycles. No external cohort was available "
            f"{cite('tripod')}. The models are a "
            f"research instrument for prioritizing exposures for the Mendelian randomization phase, "
            f"not a clinical screening tool, and have not been evaluated prospectively. Analyses used "
            f"Python 3.11–3.13 with scikit-learn, imbalanced-learn, LightGBM, XGBoost, CatBoost and SHAP, "
            f"one seed fixed throughout. Descriptive statistics used the NHANES complex-survey design "
            f"with Taylor-series linearization and domain estimation, the fasting-subsample weight "
            f"divided by the number of pooled cycles (the analytic cohort is the fasting subsample), "
            f"weighted geometric means for the metals, and design-based t and adjusted Wald F tests "
            f"{cite('nhanes_design', 'nhanes_weights')}; in 2013–2016 the metals were measured in a "
            f"separate half-sample for which no combined weight exists, so metal estimates for those "
            f"cycles are approximate; "
            f"models were fitted unweighted, the target being individual-level discrimination rather "
            f"than a population parameter. Data are public at {NHANES_URL}; code, derived metrics and "
            f"model artifacts are kept under version control and will be deposited publicly on "
            f"publication. NHANES protocols were "
            f"approved by the NCHS Research Ethics Review Board with written informed consent; this "
            f"secondary analysis required no additional approval.")

    # ==================================================================
    ds.heading(doc, "References", 1)
    for i, (_key, text) in enumerate(REFS, start=1):
        p = doc.add_paragraph()
        from docx.enum.text import WD_LINE_SPACING
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = ds.Pt(4)
        p.paragraph_format.left_indent = ds.Inches(0.3)
        p.paragraph_format.first_line_indent = ds.Inches(-0.3)
        ds._set_run(p.add_run(f"{i}. "), size=ds.Pt(9))
        ds._set_run(p.add_run(text), size=ds.Pt(9))

    path = OUT / "Manuscript_Methods_AI.docx"
    # Convert break-only paragraphs into page_break_before on the following
    # block, so a section that happens to end on a page boundary cannot leave
    # a blank page behind it.
    n_breaks = ds.normalize_page_breaks(doc)
    doc.save(path)
    print(f"Saved: {path}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} refs={len(REFS)}")


if __name__ == "__main__":
    main()
