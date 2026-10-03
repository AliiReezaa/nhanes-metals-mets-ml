"""
Rebuild the proposal's citation system.

The previous draft carried 96 markers numbered [1]-[97] and no bibliography at all, and
the numbering collided: the same number served two unrelated literatures in two different
sections ([44] was both an income/ethnicity claim in the literature review and the
inverse-variance-weighted estimator in the methods; [38] was both diagnostic cut-offs and
LD clumping). Renumbering by hand would reproduce that class of error, so numbers are not
assigned by hand anywhere. Every marker in the body is written as a mnemonic placeholder,
{{key|key}}, and this script walks the document in reading order and assigns [1], [2], ...
on first mention. A reference that is never mentioned never receives a number and never
reaches the list.

Three kinds of edit are applied, in this order:

  1. PARA_REPLACE -- paragraphs rewritten in full because at least one sentence in them
     asserted something no verifiable source supports. These are rewrites, not deletions
     of inconvenient text: each says what the literature that does exist actually shows.
     What was removed, and why, is recorded in REMOVED_CLAIMS below.
  2. SUBSTR_REPLACE -- surgical replacements inside paragraphs that are otherwise sound.
  3. CLUSTER_MAP   -- the (paragraph, cluster) -> reference-keys table for every marker
     cluster that survives unchanged. A "cluster" is a run of adjacent markers such as
     [28][29][30]; the whole run collapses to one marker.

The script refuses to write anything if the clusters it finds do not match the table
exactly, so a silently dropped or double-counted citation cannot slip through.
"""
from __future__ import annotations

import importlib.util
import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docx_final" / "Metabolic_Syndrome_Proposal_Revised_FINAL3.docx"
DST = ROOT / "docx_final" / "Metabolic_Syndrome_Proposal_Revised_FINAL4.docx"

spec = importlib.util.spec_from_file_location(
    "propref", ROOT / "scripts" / "22_proposal_references.py")
propref = importlib.util.module_from_spec(spec)
spec.loader.exec_module(propref)
REFS, format_reference, collapse = propref.REFS, propref.format_reference, propref.collapse

MARKER = re.compile(r"\[\s*\d+(?:\s*[-–,]\s*\d+)*\s*\]")
PLACEHOLDER = re.compile(r"\{\{([a-z0-9_|]+)\}\}")


# ----------------------------------------------------------------------
# 1. Paragraphs rewritten in full
# ----------------------------------------------------------------------
REMOVED_CLAIMS = """
Claims removed because no real, verifiable article was found to support them. Each was
searched for in PubMed and none of the candidate papers matched; rather than attach a
citation that does not say what the sentence claims, the sentence was rewritten.

  para 11  "in a study of peri-menopausal women, there was no significant difference in
           blood lead between those with and without MetS"
  para 11  a second, redundant attribution of the KNHANES lead-MetS finding to Rhee, which
           the preceding sentence already covers
  para 12  "Some research in the US observed higher blood lead and cadmium with LOWER MetS
           and abdominal obesity prevalence" -- the verified US/NHANES analyses report
           associations that differ in direction and magnitude from Asian-population
           findings, not a protective effect
  para 25  marital status as a risk/protective/null factor for MetS
  para 25  household size as a proxy for socioeconomic or cultural risk conditions
  para 29  "People living with HIV/AIDS ... their MetS risk rises at BMI levels that would
           look unremarkable in the general population"
  para 34  "wine being less harmful than beer or spirits"
  para 38  the same inverse-association claim as para 12
  para 70  "lead from the same 13-SNP instrument used in prior lead-cardiovascular
           disease MR work" -- no such instrument paper could be located
"""

PARA_REPLACE: dict[int, str] = {

    11: (
        "But none of this resolves the question of how strong the blood lead–metabolic "
        "syndrome link is. Several Korean studies have shown a clear association, where "
        "higher blood lead is associated with higher odds of MetS, higher triglycerides, "
        "and larger waist circumference {{me_rhee|me_leekim_2012|me_moon}}. Other studies "
        "do not see the same pattern. Findings vary across samples and age bands, and the "
        "strength of the lead association is not consistent from one study to the next, "
        "which is why the pooled estimates carry substantial heterogeneity "
        "{{me_hasani|me_xu_four|me_parkoh}}. Some data indicate that lead affects specific "
        "components of MetS — blood pressure and lipids in particular — more than it "
        "predicts the syndrome as a coherent whole {{pb_prehtn|me_park_pb}}."
    ),

    12: (
        "The picture is no better for the other metals. Large cross-sectional studies and "
        "meta-analyses report inverse or null relationships for lead, cadmium and mercury "
        "in particular populations, so the direction of association is not settled even "
        "for the two metals with the most consistent overall evidence "
        "{{me_lo_mn|me_lee_hg|me_xu_four}}. Analyses of United States populations have "
        "reported associations that differ from those in Asian cohorts in both direction "
        "and magnitude, which is a harder problem than a simple difference in effect size "
        "{{me_wang_mix|me_bulka}}. Manganese is the least clear of the group, with higher "
        "risk seen in some studies and only in certain age groups, and not seen at all in "
        "others {{me_wong_mn|me_lo_mn}}."
    ),

    25: (
        "Household-level social variables are harder to pin down than income or education. "
        "The evidence base for factors such as marital status and household composition is "
        "thin and inconsistent, and this proposal therefore carries them as descriptive "
        "covariates rather than as hypothesised risk factors. What does replicate is that "
        "where people place themselves on a social ladder tracks metabolic risk over and "
        "above their measured income and education, and does so most strongly in women "
        "{{ses_adler|ses_manuck}}."
    ),

    29: (
        "Which anthropometric measure predicts MetS best is still an open question. Some "
        "studies find BMI and hip circumference (or their derived indices) roughly "
        "equivalent; others find central obesity measures clearly superior, particularly "
        "in men and certain ethnic groups {{an_ashwell|an_bergman|an_snijder}}. Cut-off "
        "values complicate things further, since the optimal thresholds for BMI and hip "
        "circumference shift by age, sex, and ethnicity {{an_ashwell|an_who_asian}}. The "
        "point applies across populations and not only across individuals: thresholds "
        "validated in European-ancestry samples misclassify cardiometabolic risk when "
        "applied unchanged to Asian populations, where risk begins to rise at a lower body "
        "mass index {{an_who_asian}}."
    ),

    34: (
        "Alcohol is the messiest variable of all. Plenty of studies find no significant "
        "alcohol–MetS link {{ls_sun_alcohol}}. Others report that heavy or frequent "
        "drinking raises risk, while light-to-moderate drinking looks neutral or even "
        "protective {{ls_sun_alcohol}}. Differences in drinking pattern, in the sex "
        "composition of the samples, and in how the reference group is defined plausibly "
        "account for much of that disagreement, which makes population-level summaries of "
        "alcohol and MetS difficult to interpret {{ls_sun_alcohol|ls_slagter}}."
    ),

    38: (
        "Several large cross-sectional studies and meta-analyses report inverse or null "
        "associations for lead, cadmium, and mercury in some populations, so the direction "
        "of association is not settled {{me_lo_mn|me_lee_hg|me_xu_four}}. United States "
        "analyses have reported associations differing from those in Asian populations in "
        "direction as well as magnitude, rather than weaker versions of the same effect "
        "{{me_wang_mix|me_bulka}}. Manganese is the least settled of the group: increased "
        "risk shows up only in specific age brackets in some studies, and not at all in "
        "others {{me_wong_mn|me_lo_mn}}."
    ),
}


# ----------------------------------------------------------------------
# 2. Surgical replacements inside otherwise sound paragraphs
# ----------------------------------------------------------------------
SUBSTR_REPLACE: dict[int, list[tuple[str, str]]] = {
    70: [(
        "and lead from the same 13-SNP instrument used in prior lead–cardiovascular "
        "disease MR work [95].",
        "and lead from the same genome-wide meta-analyses of blood and erythrocyte metal "
        "concentrations, no larger dedicated blood-lead GWAS being available at the time "
        "of writing {{gw_ng_metals|gw_moksnes}}.",
    )],
    78: [(
        "is biologically plausible and needs to be ruled out rather than assumed away.",
        "is biologically plausible and needs to be ruled out rather than assumed away "
        "{{mr_steiger}}. The Mendelian randomization analyses will be reported in "
        "accordance with the STROBE-MR guideline {{mr_strobe}}.",
    )],
}


# ----------------------------------------------------------------------
# 3. Cluster -> reference keys, for markers that survive unchanged
# ----------------------------------------------------------------------
CLUSTER_MAP: dict[tuple[int, int], list[str]] = {
    # -- State of the problem: lead exposure and its effects
    (3, 0): ["pb_whelan", "pb_rinsky"],
    (4, 0): ["pb_needleman", "pb_lanphear_iq"],
    (4, 1): ["pb_lanphear_iq", "pb_navas", "pb_needleman"],
    (4, 2): ["pb_lanphear_iq", "pb_reuben"],

    # -- The other metals, introduced
    (5, 0): ["me_hasani", "me_xu_four", "me_lu_cd", "me_leekim_2012", "me_eom_hg"],
    (5, 1): ["me_yuan_se", "me_steinbrenner_se", "me_weening_se"],
    (5, 2): ["me_wong_mn", "me_lo_mn"],

    # -- Metabolic syndrome: what it is, what it costs
    (6, 0): ["mets_alberti2009", "mets_grundy2005", "mets_isomaa", "mets_mottillo",
             "mets_gami"],
    (7, 0): ["mets_eckel", "ls_zhang_pa", "ls_ford_sed"],
    (8, 0): ["mets_saklayen", "mets_mottillo", "mets_hirode"],

    # -- Lead and MetS: mechanisms
    (9, 0): ["me_rhee"],
    (9, 1): ["me_moon", "me_park_pb"],
    (9, 2): ["pb_navas", "pb_perlstein"],
    (9, 3): ["me_choi_tsh"],
    (10, 0): ["me_tinkov_cd", "me_steinbrenner_se"],
    (10, 1): ["me_bulka", "me_wang_mix", "me_ngu"],

    # -- Machine learning with SHAP
    (13, 0): ["ml_huang_htn", "ml_wang_htn", "ml_yan_ad", "ml_liu_dr"],
    (14, 0): ["ml_lundberg", "ml_wang_htn", "ml_wan_dr"],
    (14, 1): ["ml_he_htn", "ml_huang_htn"],
    (15, 0): ["ml_chen_pi", "ml_liang_covid"],

    # -- Mendelian randomization, introduced
    (16, 0): ["mr_dsmith2003", "mr_davies2018", "mr_ference2012"],

    # -- Literature review: age and sex
    (19, 0): ["as_hildrum", "as_moore", "mets_beltran", "mets_hirode"],
    (19, 1): ["as_pradhan", "as_carr"],
    (19, 2): ["as_regitz", "ses_manuck"],
    (20, 0): ["as_pradhan", "as_moore", "as_regitz"],
    (20, 1): ["as_carr", "as_pradhan"],
    (20, 2): ["as_hildrum", "as_kuk"],
    (21, 0): ["as_kuk", "as_hildrum"],
    (21, 1): ["as_kuk", "as_pradhan"],
    (21, 2): ["mets_alberti2009", "an_who_asian"],
    (22, 0): ["as_pradhan", "as_regitz", "as_kuk", "as_moore", "as_hildrum"],

    # -- Literature review: ethnicity and socioeconomic position
    (23, 0): ["as_moore"],
    (23, 1): ["as_moore"],
    (23, 2): ["as_moore", "ses_loucks"],
    (24, 0): ["ses_loucks"],
    (24, 1): ["ses_loucks", "as_moore"],
    (24, 2): ["ses_manuck", "ses_adler"],
    (26, 0): ["as_moore", "ses_loucks", "ses_manuck"],

    # -- Literature review: anthropometry
    (27, 0): ["an_ashwell", "mets_alberti2009"],
    (27, 1): ["an_ashwell"],
    (28, 0): ["an_ashwell", "an_bergman"],
    (28, 1): ["an_snijder", "an_manolopoulos"],
    (30, 0): ["an_snijder", "an_manolopoulos", "an_who_asian"],

    # -- Literature review: lifestyle
    (32, 0): ["ls_zhang_pa", "ls_edwardson", "ls_ford_sed"],
    (32, 1): ["ls_zhang_pa", "ls_ford_sed"],
    (33, 0): ["ls_slagter"],
    (33, 1): ["ls_sun_smoke", "ls_slagter"],
    (33, 2): ["ls_sun_smoke", "ls_slagter"],
    (35, 0): ["ls_zhang_pa", "ls_sun_smoke", "ls_sun_alcohol"],

    # -- Literature review: metals and MetS
    (37, 0): ["me_hasani", "me_xu_four", "me_lu_cd", "me_leekim_2012", "me_eom_hg",
              "me_chung_hg"],
    (37, 1): ["me_yuan_se", "me_steinbrenner_se", "me_ngu"],
    (37, 2): ["me_rhee", "me_park_pb", "me_parkoh", "me_leekim_2012"],
    (39, 0): ["me_tinkov_cd", "me_steinbrenner_se", "me_pang"],
    (39, 1): ["me_bulka", "me_wang_mix", "me_ngu"],
    (39, 2): ["me_chung_hg", "me_weening_se"],
    (40, 0): ["me_hasani", "me_xu_four", "me_yao_ml"],

    # -- Methods: genetic instrument sources
    (70, 0): ["gw_ng_metals", "gw_moksnes"],
    (70, 1): ["gw_borne_cd"],
    (70, 2): ["gw_cornelis_se"],
    # the lead-instrument clause is handled by SUBSTR_REPLACE above, so the marker it
    # carried is already gone by the time clusters are counted here
    (70, 3): ["gw_dack_hg"],
    (70, 4): ["gw_giant", "gw_yengo"],
    (70, 5): ["gw_gscan"],
    (70, 6): ["gw_doherty"],

    # -- Methods: instrument selection and harmonisation
    (73, 0): ["mr_weakinst", "mr_twosample"],
    (73, 1): ["mr_mrbase", "mr_davies2018"],
    (73, 2): ["mr_phenoscanner"],
    (73, 3): ["mr_mrbase", "mr_lawlor2008"],

    # -- Methods: estimators
    (75, 0): ["mr_ivw"],
    (75, 1): ["mr_ivw"],
    (75, 2): ["mr_egger"],
    (75, 3): ["mr_wmedian", "mr_wmode"],

    # -- Methods: heterogeneity and pleiotropy diagnostics
    (76, 0): ["mr_ruckerq"],
    (76, 1): ["mr_isq"],
    (76, 2): ["mr_egger"],
    (77, 0): ["mr_presso"],
    (77, 1): ["mr_radial"],
    (77, 2): ["mr_mrbase"],
    (78, 0): ["mr_lasso", "mr_cml", "mr_contam"],
}


# ----------------------------------------------------------------------
# Document surgery
# ----------------------------------------------------------------------
def set_paragraph_text(p, text: str) -> None:
    """Replace a paragraph's text, keeping the first run's character formatting.

    Safe only where the paragraph is formatted uniformly; check_uniform() enforces that
    before any paragraph is touched.
    """
    if not p.runs:
        p.add_run("")
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def check_uniform(p, idx: int) -> None:
    flags = {(r.bold, r.italic, r.underline, r.font.name,
              r.font.size, str(r.font.color.rgb) if r.font.color and
              r.font.color.type is not None else None)
             for r in p.runs if r.text}
    if len(flags) > 1:
        raise SystemExit(
            f"paragraph {idx} mixes character formatting across {len(flags)} run styles; "
            "rewriting it would flatten that formatting. Handle it by hand.")


def clusters_in(text: str) -> list[tuple[int, int]]:
    """Spans of adjacent citation markers, e.g. '[28][29][30]' -> one span."""
    ms = list(MARKER.finditer(text))
    if not ms:
        return []
    out, start, end = [], ms[0].start(), ms[0].end()
    for prev, cur in zip(ms, ms[1:]):
        if text[prev.end():cur.start()].strip() == "":
            end = cur.end()
        else:
            out.append((start, end))
            start, end = cur.start(), cur.end()
    out.append((start, end))
    return out


def main() -> int:
    doc = Document(SRC)
    paras = doc.paragraphs

    unknown = {k for keys in list(CLUSTER_MAP.values()) for k in keys} - set(REFS)
    for text in list(PARA_REPLACE.values()) + [n for v in SUBSTR_REPLACE.values()
                                               for _, n in v]:
        unknown |= {k for m in PLACEHOLDER.finditer(text)
                    for k in m.group(1).split("|")} - set(REFS)
    if unknown:
        raise SystemExit(f"unknown reference keys: {sorted(unknown)}")

    # --- 1. full-paragraph rewrites
    for idx, new_text in PARA_REPLACE.items():
        check_uniform(paras[idx], idx)
        set_paragraph_text(paras[idx], new_text)

    # --- 2. surgical replacements
    for idx, pairs in SUBSTR_REPLACE.items():
        check_uniform(paras[idx], idx)
        text = paras[idx].text
        for old, new in pairs:
            if old not in text:
                raise SystemExit(f"paragraph {idx}: substring not found: {old!r}")
            text = text.replace(old, new, 1)
        set_paragraph_text(paras[idx], text)

    # --- 3. cluster replacement everywhere else
    expected = {i for i, _ in CLUSTER_MAP}
    for idx, p in enumerate(paras):
        if idx in PARA_REPLACE:
            continue
        text = p.text
        spans = clusters_in(text)
        if not spans:
            if idx in expected:
                raise SystemExit(f"paragraph {idx} was expected to hold markers, found none")
            continue
        mapped = [k for k in CLUSTER_MAP if k[0] == idx]
        if len(mapped) != len(spans):
            raise SystemExit(
                f"paragraph {idx}: found {len(spans)} marker cluster(s) but the table "
                f"defines {len(mapped)}")
        check_uniform(p, idx)
        for c, (s, e) in reversed(list(enumerate(spans))):
            keys = CLUSTER_MAP[(idx, c)]
            text = text[:s] + "{{" + "|".join(keys) + "}}" + text[e:]
        set_paragraph_text(p, text)

    leftover = [i for i, p in enumerate(doc.paragraphs) if MARKER.search(p.text)]
    if leftover:
        raise SystemExit(f"old-style markers survive in paragraphs {leftover}")

    # --- 4. number in first-mention order and substitute
    order: list[str] = []

    def number_of(key: str) -> int:
        if key not in order:
            order.append(key)
        return order.index(key) + 1

    for p in doc.paragraphs:
        text = p.text
        if not PLACEHOLDER.search(text):
            continue

        def sub(m: re.Match) -> str:
            return "[" + collapse([number_of(k) for k in m.group(1).split("|")]) + "]"

        set_paragraph_text(p, PLACEHOLDER.sub(sub, text))

    # --- 5. the reference list
    heading_src = next(p for p in doc.paragraphs if p.text.strip() == "2. Literature Review")
    body_src = doc.paragraphs[3]

    body = doc.element.body
    sect_pr = body.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr")

    def append_like(template, text: str):
        """Clone a paragraph's formatting and add it at the end of the body.

        The clone is inserted BEFORE the trailing <w:sectPr>. Word requires the final
        section-properties element to be the last child of the body; appending after it
        produces a file that opens but whose page setup is silently wrong.
        """
        new_p = deepcopy(template._p)
        if sect_pr is not None:
            sect_pr.addprevious(new_p)
        else:
            body.append(new_p)
        from docx.text.paragraph import Paragraph
        para = Paragraph(new_p, template._parent)
        set_paragraph_text(para, text)
        return para

    refs_heading = append_like(heading_src, "References")
    # Start the list on a fresh page rather than partway down whatever page the methods
    # happen to end on. page_break_before is used instead of an inserted break character
    # so no empty paragraph is left behind if the text above later grows or shrinks.
    refs_heading.paragraph_format.page_break_before = True
    refs_heading.paragraph_format.keep_with_next = True
    for n, key in enumerate(order, start=1):
        append_like(body_src, format_reference(key, n))

    DST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DST)

    unused = [k for k in REFS if k not in order]
    print(f"wrote {DST.relative_to(ROOT)}")
    print(f"marker clusters rewritten : {len(CLUSTER_MAP) + sum(len(PLACEHOLDER.findall(t)) for t in PARA_REPLACE.values()) + sum(len(PLACEHOLDER.findall(n)) for v in SUBSTR_REPLACE.values() for _, n in v)}")
    print(f"references cited          : {len(order)}")
    print(f"references in bank, unused: {len(unused)}")
    print(f"paragraphs rewritten      : {sorted(PARA_REPLACE)}")
    print(f"paragraphs edited in place: {sorted(SUBSTR_REPLACE)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
