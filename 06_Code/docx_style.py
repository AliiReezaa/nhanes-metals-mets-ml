"""
Shared manuscript-docx styling helpers (journal submission format).

Conventions applied:
  * Times New Roman 12 pt, double-spaced body text -- the default submission
    format for the large majority of clinical/epidemiology journals.
  * Numbered section headings in bold roman, not Word's built-in blue Heading
    styles (which most journals strip anyway and which look wrong in print).
  * Three-line ("booktabs") tables: rule above the header, rule below the
    header, rule below the last row, and no vertical rules at all.
  * Table captions ABOVE the table, figure captions BELOW the figure, both
    single-spaced in 10 pt, with a bold "Table N."/"Figure N." lead-in.
"""
from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(12)
CAPTION_SIZE = Pt(10)
TABLE_SIZE = Pt(9)


# ----------------------------------------------------------------------
# Document setup
# ----------------------------------------------------------------------
def new_document(margin_in: float = 1.0, line_spacing: str = "double") -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(margin_in)
        section.bottom_margin = Inches(margin_in)
        section.left_margin = Inches(margin_in)
        section.right_margin = Inches(margin_in)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = (WD_LINE_SPACING.DOUBLE if line_spacing == "double"
                            else WD_LINE_SPACING.ONE_POINT_FIVE)
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    return doc


def _set_run(run, *, size=None, bold=False, italic=False, font=BODY_FONT, color=None):
    run.font.name = font
    run.font.size = size or BODY_SIZE
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    return run


# ----------------------------------------------------------------------
# Text blocks
# ----------------------------------------------------------------------
def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    _set_run(p.add_run(text), size=Pt(14), bold=True)
    return p


def subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    _set_run(p.add_run(text), size=Pt(11), italic=True)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.keep_with_next = True
    size = {1: Pt(12.5), 2: Pt(12), 3: Pt(12)}[level]
    _set_run(p.add_run(text), size=size, bold=True, italic=(level == 3))
    return p


def para(doc, text, first_line_indent=True, space_after=0):
    """Body paragraph. Supports **bold** and *italic* inline markers.
    Inherits the document's line-spacing default set in new_document()."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    if first_line_indent:
        p.paragraph_format.first_line_indent = Inches(0.3)
    _add_marked_runs(p, text)
    return p


def _add_marked_runs(p, text):
    import re
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            _set_run(p.add_run(tok[2:-2]), bold=True)
        elif tok.startswith("*") and tok.endswith("*"):
            _set_run(p.add_run(tok[1:-1]), italic=True)
        else:
            _set_run(p.add_run(tok))


def caption(doc, label, text, above=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(4 if above else 6)
    p.paragraph_format.space_after = Pt(8 if above else 14)
    p.paragraph_format.keep_with_next = above
    _set_run(p.add_run(label + " "), size=CAPTION_SIZE, bold=True)
    _add_caption_runs(p, text)
    return p


def _add_caption_runs(p, text):
    import re
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            _set_run(p.add_run(tok[2:-2]), size=CAPTION_SIZE, bold=True)
        elif tok.startswith("*") and tok.endswith("*"):
            _set_run(p.add_run(tok[1:-1]), size=CAPTION_SIZE, italic=True)
        else:
            _set_run(p.add_run(tok), size=CAPTION_SIZE)


def figure(doc, image_path, width_in=6.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(width_in))
    return p


def page_break(doc):
    """Mark where the next block should start on a fresh page.

    This inserts a paragraph carrying a page-break character, which is convenient while
    the document is being written because the next block does not exist yet. It has a
    failure mode, though: if the preceding content happens to end exactly at a page
    boundary, the break paragraph is pushed onto a page of its own and the reader gets a
    blank page. normalize_page_breaks() is therefore run once at the end of every
    document build to convert these into page_break_before on the following block, which
    cannot produce an empty page.
    """
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def _is_break_only(p_el) -> bool:
    """A paragraph whose entire content is one or more page-break characters."""
    if p_el.tag != qn("w:p"):
        return False
    if "".join(t.text or "" for t in p_el.iter(qn("w:t"))).strip():
        return False
    breaks = [b for b in p_el.iter(qn("w:br"))
              if b.get(qn("w:type")) == "page"]
    return bool(breaks)


def _set_page_break_before(p_el) -> None:
    pPr = p_el.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_el.insert(0, pPr)
    existing = pPr.find(qn("w:pageBreakBefore"))
    if existing is None:
        existing = OxmlElement("w:pageBreakBefore")
        pPr.append(existing)
    existing.set(qn("w:val"), "true")


def normalize_page_breaks(doc) -> int:
    """Turn every break-only paragraph into page_break_before on the block after it.

    Returns how many were converted. A break-only paragraph at the very end of the body
    is simply deleted: it can only ever produce a trailing blank page.
    """
    body = doc.element.body
    converted = 0
    for p_el in list(body.iter(qn("w:p"))):
        if not _is_break_only(p_el):
            continue
        parent = p_el.getparent()
        target = p_el.getnext()
        # step over any further break-only paragraphs stacked together
        while target is not None and _is_break_only(target):
            nxt = target.getnext()
            target.getparent().remove(target)
            target = nxt
        if target is None or target.tag == qn("w:sectPr"):
            parent.remove(p_el)
            converted += 1
            continue
        if target.tag == qn("w:tbl"):
            first_p = next(target.iter(qn("w:p")), None)
            if first_p is not None:
                _set_page_break_before(first_p)
        else:
            _set_page_break_before(target)
        parent.remove(p_el)
        converted += 1
    return converted


# ----------------------------------------------------------------------
# Tables (three-line academic style)
# ----------------------------------------------------------------------
def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "bottom", "left", "right"):
        spec = kwargs.get(edge)
        tag = "w:{}".format(edge)
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec.get("sz", 8)))
            el.set(qn("w:color"), spec.get("color", "000000"))
            el.set(qn("w:space"), "0")


def _repeat_header(row):
    """Mark a table row as a header row so Word repeats it on every page the
    table spans."""
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _no_split(row):
    """Stop a single row from being broken across a page boundary."""
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def add_table(doc, headers, rows, col_widths_in=None, font_size=TABLE_SIZE,
              align_right_from=1, header_wrap=True):
    """Three-line table: rule above header, below header, below final row."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _repeat_header(t.rows[0])
    _no_split(t.rows[0])

    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i < align_right_from else WD_ALIGN_PARAGRAPH.CENTER
        _set_run(p.add_run(str(h)), size=font_size, bold=True)
        _set_cell_border(hdr[i], top={"sz": 12}, bottom={"sz": 8}, left=None, right=None)

    for r_i, row in enumerate(rows):
        trow = t.add_row()
        _no_split(trow)
        cells = trow.cells
        last = (r_i == len(rows) - 1)
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.space_before = Pt(1.5)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i < align_right_from else WD_ALIGN_PARAGRAPH.CENTER
            txt = str(val)
            bold = txt.startswith("**") and txt.endswith("**")
            if bold:
                txt = txt[2:-2]
            _set_run(p.add_run(txt), size=font_size, bold=bold)
            _set_cell_border(cells[i], top=None,
                             bottom={"sz": 12} if last else None, left=None, right=None)

    if col_widths_in:
        for row in t.rows:
            for i, w in enumerate(col_widths_in):
                row.cells[i].width = Inches(w)
    return t


def table_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(14)
    _add_caption_runs(p, text)
    for r in p.runs:
        r.font.size = Pt(9)
    return p
